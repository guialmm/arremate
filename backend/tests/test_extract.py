import io
import zipfile

import docx
import pytest

from app.ingest.extract import FormatoNaoSuportado, detectar_formato, extrair, limpar
from tests.conftest import make_pdf, make_zip

HABILITACAO = (
    "9.1. Para habilitacao, o licitante devera apresentar atestado de capacidade tecnica "
    "comprovando fornecimento de pelo menos 50% do quantitativo licitado."
)
PRAZO = "6.2. O prazo de entrega sera de 15 dias corridos contados da ordem de fornecimento."


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_pdf_pages_keep_their_numbers():
    paginas = extrair(make_pdf([HABILITACAO, PRAZO]), "edital.pdf")
    assert [p.pagina_no_arquivo for p in paginas] == [1, 2]
    assert "atestado de capacidade tecnica" in paginas[0].texto
    assert "15 dias corridos" in paginas[1].texto
    assert paginas[0].arquivo is None


def test_scanned_pdf_is_unsupported():
    with pytest.raises(FormatoNaoSuportado, match="digitalizado"):
        extrair(make_pdf(["", "1"]), "scan.pdf")


def test_broken_pdf_is_unsupported():
    with pytest.raises(FormatoNaoSuportado, match="PDF ilegível"):
        extrair(b"%PDF-1.4 truncated garbage", "quebrado.pdf")


def test_format_comes_from_content_not_from_the_name():
    assert detectar_formato(make_pdf([HABILITACAO]), "edital.doc") == "pdf"
    assert detectar_formato(make_zip({"a.pdf": b"x"}), "edital.pdf") == "zip"
    assert detectar_formato(make_docx(["x"]), "edital.zip") == "docx"
    assert detectar_formato(b"\xd0\xcf\x11\xe0legacy", "edital.pdf") == "doc"
    assert detectar_formato(b"Rar!\x1a\x07", None) == "rar"


def test_legacy_formats_are_unsupported():
    with pytest.raises(FormatoNaoSuportado, match="formato doc"):
        extrair(b"\xd0\xcf\x11\xe0" + b"\x00" * 100, "edital.doc")
    with pytest.raises(FormatoNaoSuportado, match="formato rar"):
        extrair(b"Rar!\x1a\x07\x00", "edital.rar")


def test_docx_is_split_into_page_sized_chunks_with_tables():
    paragraphs = [f"Clausula {i}. " + "texto " * 100 for i in range(20)]
    paginas = extrair(make_docx(paragraphs, [["Documento", "Exigido"], ["Balanco", "Sim"]]))
    assert len(paginas) > 3
    assert all(len(p.texto) <= 3000 + 700 for p in paginas)
    assert "Clausula 0." in paginas[0].texto
    assert "Balanco | Sim" in paginas[-1].texto


def test_zip_reads_every_file_with_the_edital_first():
    conteudo = make_zip(
        {
            "Anexo II - Modelo de Proposta.pdf": make_pdf([PRAZO]),
            "Anexo I - Termo de Referencia.docx": make_docx([HABILITACAO]),
            "Edital Pregao 24-2026.pdf": make_pdf([HABILITACAO, PRAZO]),
            "__MACOSX/._lixo.pdf": b"junk",
            "planilha.xlsx": b"PK\x03\x04not really",
        }
    )
    paginas = extrair(conteudo, "edital.zip")
    assert [p.arquivo for p in paginas] == [
        "Edital Pregao 24-2026.pdf",
        "Edital Pregao 24-2026.pdf",
        "Anexo I - Termo de Referencia.docx",
        "Anexo II - Modelo de Proposta.pdf",
    ]
    assert [p.pagina_no_arquivo for p in paginas] == [1, 2, 1, 1]


def test_zip_annexes_follow_natural_order():
    conteudo = make_zip({f"Anexo {n}.pdf": make_pdf([HABILITACAO]) for n in (10, 2, 1)})
    assert [p.arquivo for p in extrair(conteudo)] == ["Anexo 1.pdf", "Anexo 2.pdf", "Anexo 10.pdf"]


def test_zip_names_from_windows_keep_their_accents(monkeypatch):
    # Windows' built-in zipper writes OEM code page 850 bytes without the UTF-8 flag;
    # Python would write UTF-8 and set the flag, so force the Windows behaviour.
    monkeypatch.setattr(
        zipfile.ZipInfo,
        "_encodeFilenameFlags",
        lambda self: (self.filename.encode("cp850"), self.flag_bits & ~0x800),
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("TERMO DE REFERÊNCIA.pdf", make_pdf([HABILITACAO]))
    monkeypatch.undo()
    assert extrair(buf.getvalue())[0].arquivo == "TERMO DE REFERÊNCIA.pdf"


def test_zip_without_readable_files_is_unsupported():
    with pytest.raises(FormatoNaoSuportado, match="nenhum texto"):
        extrair(make_zip({"planta.dwg": b"AC1032", "foto.jpg": b"\xff\xd8"}))


def test_zip_bomb_is_refused(monkeypatch):
    monkeypatch.setattr("app.ingest.extract.MAX_DESCOMPACTADO", 1000)
    with pytest.raises(FormatoNaoSuportado, match="grande demais"):
        extrair(make_zip({"a.pdf": b"0" * 2000}))


def test_nested_zips_stop_at_the_depth_limit():
    interno = make_zip({"edital.pdf": make_pdf([HABILITACAO])})
    assert len(extrair(make_zip({"pacote.zip": interno}))) == 1
    with pytest.raises(FormatoNaoSuportado):
        extrair(make_zip({"a.zip": make_zip({"b.zip": interno})}))


def test_limpar_removes_nul_and_collapses_whitespace():
    assert limpar("  Edital\x00 nº  12\n\n\n\n  Anexo \t I ") == "Edital nº 12\n\nAnexo I"
